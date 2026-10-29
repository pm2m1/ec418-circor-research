"""Write-up number check: fail if the report or the slides contain a number that does not come from the registry.

Report (report/EC418_report.tex, report/tables/*.tex and submission/report_src/*.tex): registry numbers enter only through \\nm{key} / \\N... macros, so
after removing those macros, LaTeX commands' arguments that are not prose (\\cite, \\label, \\ref, \\input, \\includegraphics,
\\graphicspath, \\texttt code/identifiers, lengths), comments and the bibliography, every remaining digit string is typed by
hand and must be listed in report/number_whitelist.txt (token, reason: years, dataset facts, protocol constants...).
Every \\nm{key} must exist in report/numbers.json.
Slides (report/EC418_slides.pptx and submission/EC418_Assignment2_Presentation.pptx, slide text and speaker notes): every digit string must occur in some registry value
(report/numbers.json) or be whitelisted.
Exit code 1 on any violation.
usage: .venv/bin/python scripts/p4_check_numbers.py [--report] [--slides]   (default: both, slides only if the file exists)
"""
import glob
import json
import os
import re
import sys

N = json.load(open("report/numbers.json"))
WL = {}
for line in open("report/number_whitelist.txt"):
    line = line.split("#")[0].strip()
    if line:
        tok, why = [x.strip() for x in line.split("|", 1)]
        WL[tok] = why
# a free-standing number: not part of an identifier such as B0, S1, H6a, E1b, v1, A_scat12 or a file name
NUM = re.compile(r"(?<![A-Za-z_\d.])\d+(?:[.,]\d+)*(?![A-Za-z_\d])")
REG = {t for v in N.values() for t in NUM.findall(str(v["value"]))}
bad = []


def strip_tex(s):
    s = re.sub(r"(?<!\\)%.*", "", s)
    s = re.sub(r"\\begin\{thebibliography\}.*?\\end\{thebibliography\}", "", s, flags=re.S)
    keys = re.findall(r"\\nm\{([^}]*)\}", s)
    for k in keys:
        if k not in N:
            bad.append(f"report: unknown registry key \\nm{{{k}}}")
    s = re.sub(r"\\nm\{[^}]*\}", " ", s)
    s = re.sub(r"\\N[A-Z][A-Za-z]*", " ", s)
    s = re.sub(r"\\(cite|label|ref|input|includegraphics|graphicspath|texttt|url|href|bibitem|documentclass|usepackage)"
               r"(\[[^\]]*\])?\{([^{}]|\{[^{}]*\})*\}", " ", s)
    s = re.sub(r"\{\{[^}]*\}\{[^}]*\}\}", " ", s)  # \graphicspath leftovers
    s = re.sub(r"(width|itemsep)\s*=?\s*[\d.]+\\?\w*", " ", s)
    s = re.sub(r"\\setlength\\itemsep\{[^}]*\}", " ", s)
    s = re.sub(r"\[\s*[\d.]+\\(linewidth|textwidth)\]", " ", s)
    s = re.sub(r"\{[\d.]+\\(linewidth|textwidth)\}", " ", s)  # \resizebox{0.85\textwidth}
    s = re.sub(r"\{@\{\}[lcr@{}]+\}", " ", s)
    s = re.sub(r"\\parbox\{[\d.]+\\linewidth\}", " ", s)
    s = re.sub(r"\\newcommand\{\\\w+\}\[\d\]", " ", s)
    s = re.sub(r"#\d", " ", s)
    return s


def check_report():
    files = (["report/EC418_report.tex"] + sorted(glob.glob("report/tables/*.tex"))
             + sorted(glob.glob("submission/report_src/*.tex")))
    for f in files:
        for i, line in enumerate(strip_tex(open(f).read()).split("\n"), 1):
            for t in NUM.findall(line):
                if t not in WL:
                    bad.append(f"{f}: typed number '{t}' in: {line.strip()[:110]}")
    print(f"report: {len(files)} files checked")


def check_slides(path):
    from pptx import Presentation
    P = Presentation(path)
    for i, sl in enumerate(P.slides, 1):
        texts = [p.text for sh in sl.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]
        texts += [c.text for sh in sl.shapes if sh.has_table for row in sh.table.rows for c in row.cells]
        if sl.has_notes_slide:
            texts.append(sl.notes_slide.notes_text_frame.text)
        for t in NUM.findall(" ".join(texts)):
            if t not in REG and t not in WL:
                bad.append(f"{path} slide {i}: number '{t}' not in the registry or whitelist")
    print(f"{path}: {len(P.slides)} slides checked")


args = sys.argv[1:] or ["--report", "--slides"]
if "--report" in args:
    check_report()
for deck in ["report/EC418_slides.pptx", "submission/EC418_Assignment2_Presentation.pptx"]:
    if "--slides" in args and os.path.exists(deck):
        check_slides(deck)
if bad:
    print("\n".join(bad))
    print(f"FAIL: {len(bad)} problem(s)")
    sys.exit(1)
print("number check passed")
