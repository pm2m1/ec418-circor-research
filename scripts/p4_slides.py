"""Write-up slides: report/EC418_slides.pptx (python-pptx, 16:9), numbers only from report/numbers.json via v(key).

One message per slide, body text >= 18 pt, speaker notes on every slide. Figures: report/figures/slide_*.png
(scripts/p4_figures.py) and the frozen results/figures/fig7_packet_frequency_responses.png. The WEKA slide has placeholders for
the Explorer screenshots in weka_demo/screenshots/ (taken by the authors); a screenshot that exists is inserted instead.
Check with scripts/p4_check_numbers.py; render with LibreOffice to PNG.
usage: .venv/bin/python scripts/p4_slides.py
"""
import json
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

N = json.load(open("report/numbers.json"))
INK, MUTED, ACCENT, LIGHT = RGBColor(0x1F, 0x1F, 0x1E), RGBColor(0x6B, 0x6A, 0x64), RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xED, 0xED, 0xE8)
FIG = "report/figures"


def v(key):
    assert key in N, f"unknown registry key {key}"
    return N[key]["value"]


def rng(key):  # family range, or the single value when all representations agree
    lo, hi = v(key + "Min"), v(key + "Max")
    return lo if lo == hi else f"{lo} to {hi}"


P = Presentation()
P.slide_width, P.slide_height = Inches(13.333), Inches(7.5)
BLANK = P.slide_layouts[6]


def slide(title, notes):
    s = P.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(34), True, INK
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.3), Inches(1.2), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    s.notes_slide.notes_text_frame.text = notes
    return s


def bullets(s, items, left=0.6, top=1.6, width=12.1, height=5.4, size=22):
    tb = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        sub = it.startswith("  ")
        parts = it.strip().split("|")  # "label|rest": bold label; "a|b|c": bold middle
        bold = [True, False] if len(parts) == 2 else [False, True, False] if len(parts) == 3 else [False]
        for j, (txt, b) in enumerate(zip(parts, bold)):
            r = p.add_run()
            r.text = (("– " if sub else "• ") if j == 0 else "") + txt
            r.font.size, r.font.color.rgb, r.font.bold = Pt(size - 3 if sub else size), (MUTED if sub else INK), b
        p.space_after = Pt(8)
        p.level = 1 if sub else 0
    return tb


def image(s, path, left, top, max_w, max_h):
    w, h = Image.open(path).size
    sc = min(max_w / w, max_h / h)
    W, Hh = w * sc, h * sc
    s.shapes.add_picture(path, Inches(left + (max_w - W) / 2), Inches(top + (max_h - Hh) / 2), Inches(W), Inches(Hh))


def box(s, text, left, top, width, height, size=18):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = LIGHT
    sh.line.color.rgb = MUTED
    tf = sh.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = text
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.size, tf.paragraphs[0].font.color.rgb = Pt(size), MUTED


# 1 title
s = slide("Timing, Not Transform", "This talk asks which part of a classical heart-sound pipeline carries the murmur information. "
          "Every number on these slides is generated from saved result tables by a script, and every hypothesis was "
          "pre-registered with a fixed decision rule.")
bullets(s, ["Shortcut-controlled and externally validated murmur detection on CirCor",
            "Parth Mishra — National Institute of Technology Warangal"], top=2.3, size=26)

# 2 problem
s = slide("Problem: what actually helps murmur detection?",
          "Many murmur detectors segment each heart sound into cardiac cycles and use wavelet features. We separate the two "
          "ingredients, timing and transform, and we also ask how much a result depends on expert annotations that a deployed "
          "system would not have at test time.")
bullets(s, ["Murmurs are timed events: systolic or diastolic, between S1 and S2",
            "Question 1: does |cardiac-cycle timing| help?",
            "Question 2: does the |time–frequency transform| (wavelets vs Fourier) matter?",
            "Question 3: does a model |transfer| to an independent cohort?",
            "Caveat we test explicitly: expert segmentation at test time is an oracle"])

# 3 dataset
s = slide("Data: CirCor for development, ZCHSound for external testing",
          "CirCor is the PhysioNet 2022 Challenge training set with expert murmur labels and segmentation. We use Present versus "
          "Absent and exclude Unknown. ZCHSound is an independent Chinese paediatric cohort labelled with congenital heart "
          "disease diagnoses, not murmurs, and it was never used for training, tuning or thresholds.")
bullets(s, ["CirCor DigiScope v1.0.3: |942 patients, 3,163 recordings, 4 sites",
            "  binary task: 874 patients (179 Present, 695 Absent), 816 subjects",
            "  expert S1 / systole / S2 / diastole annotations",
            "  5 × 5 subject-grouped folds, fixed before any result",
            f"ZCHSound (Jia et al., 2024): |{v('zchNClean')} clean recordings ({v('zchNchdClean')} CHD) + {v('zchNNoisy')} noisy neonatal",
            "  labels are CHD diagnoses (VSD, ASD, PDA, PFO, NORMAL), not murmurs"])

# 4 pipeline
s = slide("Pipeline and controls",
          "Each recording is segmented, band features are computed on a cycle-aligned grid, and patient vectors go to WEKA. "
          "The controls keep the same features but destroy the timing, or use whole-signal features, or give the baseline the "
          "same span information, so each contrast isolates one ingredient.")
image(s, f"{FIG}/slide_pipeline.png", 0.4, 1.7, 12.5, 4.0)
bullets(s, ["Segmentation: |expert labels (oracle), LR-HSMM, or HSMM + learned span detector (E4, E4q)"], top=5.9, height=1.0, size=20)

# 5 wavelets
s = slide("Representations: wavelet packets, STFT bands, scattering",
          "The wavelet packet transform splits the signal into 16 equal bands; the STFT bands are matched to the same edges and time "
          "grid, so wavelet against Fourier is a like-for-like contrast. Wavelet scattering adds a modulation layer and was "
          "pre-registered as a separate test.")
bullets(s, ["WPD db4, depth 4: |16 bands of 62.5 Hz at 2 kHz (packet responses below)",
            "STFT bands matched: |same edges, time grid and dimension",
            "Scattering: |kymatio, J = 6, Q = (8, 1), first and second order",
            "Features per band: |systolic vs S1, systolic vs diastolic"], top=1.55, height=2.6, size=20)
image(s, "results/figures/fig7_packet_frequency_responses.png", 0.5, 4.05, 12.3, 3.3)

# 6 WEKA
s = slide("WEKA: classifiers and a reproducible fold",
          "All models are WEKA 3.9.7 classifiers called in batch. WEKA's own cross-validation is never used, because it splits "
          "rows and would put recordings of the same child in training and test. The command-line run of one fold reproduces "
          "the pipeline's predictions exactly, and the three Explorer screenshots will show the same fold in the GUI.")
bullets(s, ["RandomForest -I 300 -K 0 -S 1 (primary); SMO, linear PolyKernel (secondary)",
            "Subject-grouped folds from Python, never WEKA row-wise CV",
            f"Demo fold A_stft r0 f0: {v('wekaDemoNTrain')} train / {v('wekaDemoNTest')} test patients; CLI = pipeline, exactly",
            f"At 0.5: TN {v('wekaDemoTN')}, FP {v('wekaDemoFP')}, FN {v('wekaDemoFN')}, TP {v('wekaDemoTP')}"],
        top=1.6, height=2.6, size=20)
for i, (fn, lab) in enumerate([("1_preprocess.png", "Preprocess tab"), ("2_rf_options.png", "RF options dialog"),
                               ("3_result_buffer.png", "Result buffer + confusion matrix")]):
    path = f"weka_demo/screenshots/{fn}"
    if os.path.exists(path):
        image(s, path, 0.6 + i * 4.1, 4.3, 3.9, 2.8)
    else:
        box(s, f"Screenshot placeholder\n{lab}\n{fn}", 0.6 + i * 4.1, 4.3, 3.9, 2.8)

# 7 finding 1
s = slide("Finding 1: timing carries the information",
          "Keeping the features but destroying their alignment to the cardiac cycle costs about a sixth of the AUROC, in all 25 "
          "folds. The diagnostics split the gain exactly: within-span timing explains it, while the expert's choice of spans "
          "contributes nothing positive. Rows with a star use the expert annotation at test time and are diagnostics only.")
image(s, f"{FIG}/slide_decomposition.png", 0.3, 1.5, 8.4, 5.8)
bullets(s, [f"Timing destroyed (perm): |{v('hFivePermRange')}",
            f"Shifted labels: |{v('hFiveShiftRange')}",
            f"Within-span timing: |{v('decDiffWithinSpanTimingRange')}",
            f"Span choice without timing: {v('decDiffSpanSelNoTimingRange')}"], left=8.8, top=1.7, width=4.3, height=5.4, size=19)

# 8 finding 2
s = slide("Finding 2: the transform does not matter",
          "Under every labelling, wavelet packets are at most equal to matched STFT bands, and equivalence within the "
          "pre-registered margin is not shown either. Wavelet scattering does not beat STFT bands; its second-order gain is below "
          "the pre-registered 0.03.")
image(s, f"{FIG}/slide_wavelet.png", 0.3, 1.5, 8.2, 5.8)
bullets(s, [f"WPD − STFT: |{v('wvsAllRange')}",
            f"Scattering − STFT (H11a): {v('gFourH11aAScat1AStftFull')}",
            f"2nd − 1st order (H11c): {v('gFourH11cAScat12AScat1Full')}",
            "No wavelet advantage anywhere"], left=8.6, top=1.7, width=4.5, height=5.4, size=19)

# 9 automatic + shortcut
s = slide("Automatic pipeline and the shortcut check",
          "With a learned span detector the automatic pipeline matches expert labels. But its span fraction is itself predictive, "
          "so we built a shortcut-free version, E4q. E4q keeps the AUROC and beats random windows, yet its margin over a "
          "whole-signal STFT baseline with the same span information has a confidence interval that touches zero, so that "
          "advantage is not established.")
bullets(s, [f"E4 vs expert labels (H9a): |{rng('famRangeH9a')}, non-inferior",
            f"E4 vs B1 (H9b): |{rng('famRangeH9b')}",
            f"Span fraction alone predicts the class: AUROC {v('eFourShortcutPredFrac')}",
            f"E4q removes it: shortcut alone {v('gTwoAucG2SCqOnly')}; missing features {v('gTwoNofeatE4RAbsent')} % → {v('gTwoNofeatE4qAbsent')} % (Absent)",
            f"E4q vs random windows (H12b): |{v('gTwoH12bStftFull')}",
            f"E4q vs STFT + span info (H12a): |{v('gTwoH12aStftFull')} → not established"], size=21)

# 10 finding 3
s = slide(f"Finding 3: external AUROC {v('hTenA')} with the physiological order",
          "Trained once on all CirCor patients and applied unchanged, the whole-signal STFT model detects CHD on ZCHSound. The "
          "subtypes follow the pre-registered physiological order: loud shunts are detected best, the silent PFO is at chance. "
          "On the noisy neonatal set every model is at chance.")
image(s, f"{FIG}/slide_external.png", 0.3, 1.5, 8.6, 4.4)
bullets(s, [f"Clean set (H10a): |{v('hTenAFull')}",
            f"VSD − PFO (H10b): {v('hTenBvsdPfoFull')}",
            f"VSD − ASD: {v('hTenBvsdAsdFull')}",
            f"Noisy neonates: {v('zchWStftRFAllNoisy')} (chance)"], left=9.0, top=1.7, width=4.1, height=5.0, size=19)
bullets(s, ["Labels are CHD diagnoses, not murmurs; age does not explain the result (post-hoc check)"], top=6.1, height=0.9, size=18)

# 11 literature + case study
s = slide("Comparison with recent work",
          "Recent papers report much higher numbers than ours, and these two show why they are not on the same scale. One "
          "reports a balanced accuracy at a single threshold with test segments cut at expert S1 labels; the other reports "
          "snippet accuracy on annotated periods. On the same balanced-accuracy scale our expert-label model is far lower.")
bullets(s, [f"Nogueira et al. 2026: |(Se+Sp)/2 = {v('litNogueiraOverall')}, test segments start at expert S1 labels",
            f"Ours on the same scale (expert labels): {v('opE2E2StftTrueRFBacc')}; E4 {v('opE4E4StftRFBacc')}",
            f"Nie et al. 2024: {v('litNieAccuracy')} % on 5-s snippets from annotated periods",
            f"Oracle conditioning is worth up to {v('gThreeAE1bStftAuto')} AUROC in our data",
            "Kutlu & Çolak 2026 already validate CirCor models externally (CinC 2016, Yaseen)"], top=1.7, height=5.2, size=22)

# 12 limitations
s = slide("Limitations (all negatives reported)",
          "These are the results that did not work or that limit what we can claim. In particular, the advantage of the timing "
          "pipeline over a shortcut-aware STFT baseline is not established, and ZCHSound tests CHD rather than murmurs.")
bullets(s, [f"H12a not established: {v('gTwoH12aStftFull')}",
            "Noisy neonatal recordings: every model at chance",
            "Scattering (G4) negative; SMO unstable on scattering features",
            "ZCHSound labels are CHD, not murmur; PFO at chance",
            "No CirCor hidden test set; Unknown class excluded",
            "E1b, E4, S1 and Phase 3 designed after earlier results on the same folds",
            "B0's external score is an artefact (recording characteristics)"], size=21)

# 13 conclusion
s = slide("Conclusion",
          "Timing matters, the transform does not, and a simple whole-signal model transfers to an independent CHD cohort with "
          "the expected physiological ordering. Future work is quality gating for noisy recordings, learned models under the "
          "same protocol, a murmur-labelled external set, and a systematic audit of test-time expert segmentation in the literature.")
bullets(s, ["Cardiac-cycle |timing| carries the murmur information",
            "The |transform| (WPD, STFT, scattering) does not matter",
            "External transfer to paediatric CHD with the |physiological order|",
            "Next: quality gating, learned models under the same protocol, murmur-labelled external data",
            "Future: a systematic audit of test-time expert segmentation in published CirCor work"], size=24)

# 14 reproducibility
s = slide("Reproducibility",
          "Everything is versioned: protocols were tagged before the runs, results are frozen with hash manifests, and Phase 1 "
          "was reproduced on a second platform. The numbers on these slides and in the report are generated from the saved "
          "tables and checked automatically.")
bullets(s, ["Pre-registration tags: v1.0-pilot, protocol-v2-lock, e1b/diag/e4/s1-prereg; protocol-v3-prereg in the cloud-sandbox history",
            "Frozen results: results-final and phase3-results, with SHA-256 manifests",
            "Linux (Phase 1) reproduced on macOS; WEKA CLI = pipeline exactly",
            "Numbers registry: scripts/p4_numbers.py; check: scripts/p4_check_numbers.py"], size=21)

# 15 backup: hypotheses
s = slide("Backup: every pre-registered hypothesis (RF)",
          "This table summarises every pre-registered hypothesis family with the range over representations and the verdict. "
          "The full table with confidence intervals is in the report appendix.")
rows = [("H2a/b", "WPD − STFT (aligned / whole)", f"{v('famRangeH2aMin')} / {v('famRangeH2bMin')}", "not supported"),
        ("H3", "aligned − whole WPD", v("famRangeH3Min"), "supported"),
        ("H4a/b", "aligned WPD − B0 / B1", f"{v('famRangeH4aMin')} / {v('famRangeH4bMin')}", "no / supported"),
        ("H5a–d", "expert − timing destroyed", f"{v('hFivePermMin')} to {v('hFiveShiftMax')}", "supported"),
        ("H6, H8", "naive / HSMM − expert", f"{v('famRangeH8cAMin')} to {v('famRangeH6aMax')}", "not supported"),
        ("H9a/b", "E4 − expert / B1", f"{rng('famRangeH9a')} / {rng('famRangeH9b')}", "supported"),
        ("H9c", "E4 − W_stft", rng("famRangeH9c"), "stft, wpds only"),
        ("S1", "E4 − (W_stft + SC)", v("sOneWstftRange"), "possibly shortcut"),
        ("H10a/b", "ZCHSound AUROC / subtype order", f"{v('hTenA')} / {v('hTenBvsdPfo')}", "supported"),
        ("H11a–c", "scattering − STFT", f"{v('gFourH11aAScat1AStft')} to {v('gFourH11cAScat12AScat1')}", "not supported"),
        ("H12a", "E4q − (W_stft + SC_q)", rng("famRangeH12a"), "possibly shortcut"),
        ("H12b", "E4q − random windows", rng("famRangeH12b"), "supported")]
tbl = s.shapes.add_table(len(rows) + 1, 4, Inches(0.6), Inches(1.5), Inches(12.1), Inches(5.4)).table
for j, (hdr, w) in enumerate(zip(["Hyp.", "Contrast", "Δ AUROC", "Verdict"], [1.4, 4.2, 4.2, 2.3])):
    tbl.columns[j].width = Inches(w)
    tbl.cell(0, j).text = hdr
for i, r in enumerate(rows, 1):
    for j, val in enumerate(r):
        tbl.cell(i, j).text = val
for i in range(len(rows) + 1):
    for j in range(4):
        for p in tbl.cell(i, j).text_frame.paragraphs:
            p.font.size = Pt(18)
            p.space_before = p.space_after = Pt(0)
            p.font.bold = i == 0
        tbl.cell(i, j).margin_top = tbl.cell(i, j).margin_bottom = Inches(0.02)

P.save("report/EC418_slides.pptx")
print(f"slides: {len(P.slides)} -> report/EC418_slides.pptx")
