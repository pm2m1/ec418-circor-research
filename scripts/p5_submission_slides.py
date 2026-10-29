"""EC418 submission presentation: submission/EC418_Assignment2_Presentation.pptx (python-pptx, 16:9).

Numbers only from report/numbers.json via v(key) (checked by scripts/p4_check_numbers.py); every chart is a saved figure
generated from saved tables (results/figures*, report/figures). Body text >= 18 pt; 2-4 sentence speaker notes per slide.
WEKA Explorer screenshots are inserted from weka_demo/screenshots/ when present, otherwise labelled placeholders are shown.
Bracketed grey text marks metadata the author still has to add. Render check: LibreOffice headless -> PDF -> PNG.
usage: .venv/bin/python scripts/p5_submission_slides.py
"""
import json
import os

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

N = json.load(open("report/numbers.json"))
OUT = "submission/EC418_Assignment2_Presentation.pptx"
INK, MUTED, ACCENT, LIGHT = RGBColor(0x1F, 0x1F, 0x1E), RGBColor(0x6B, 0x6A, 0x64), RGBColor(0x2A, 0x78, 0xD6), RGBColor(0xED, 0xED, 0xE8)
RF, RFF, RR = "results/figures", "results/figures_final", "report/figures"


def v(key):
    assert key in N, f"unknown registry key {key}"
    return N[key]["value"]


def rng(key):
    lo, hi = v(key + "Min"), v(key + "Max")
    return lo if lo == hi else f"{lo} to {hi}"


P = Presentation()
P.slide_width, P.slide_height = Inches(13.333), Inches(7.5)
BLANK = P.slide_layouts[6]


def slide(title, notes, cite=None):
    s = P.slides.add_slide(BLANK)
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(12.1), Inches(0.9))
    tb.text_frame.word_wrap = True
    p = tb.text_frame.paragraphs[0]
    p.text = title
    p.font.size, p.font.bold, p.font.color.rgb = Pt(32), True, INK
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(1.22), Inches(1.2), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()
    if cite:
        c = s.shapes.add_textbox(Inches(0.6), Inches(6.95), Inches(12.1), Inches(0.45)).text_frame
        c.word_wrap = True
        c.paragraphs[0].text = cite
        c.paragraphs[0].font.size, c.paragraphs[0].font.color.rgb = Pt(12), MUTED
    s.notes_slide.notes_text_frame.text = notes
    return s


def bullets(s, items, left=0.6, top=1.5, width=12.1, height=5.3, size=22):
    tf = s.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height)).text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        sub = it.startswith("  ")
        parts = it.strip().split("|")  # "label|rest": bold label; "a|b|c": bold middle
        bold = [True, False] if len(parts) == 2 else [False, True, False] if len(parts) == 3 else [False]
        for j, (txt, b) in enumerate(zip(parts, bold)):
            r = p.add_run()
            r.text = (("– " if sub else "• ") if j == 0 else "") + txt
            r.font.size, r.font.bold = Pt(size - 3 if sub else size), b
            r.font.color.rgb = MUTED if sub else INK
        p.space_after = Pt(7)


def image(s, path, left, top, max_w, max_h):
    w, h = Image.open(path).size
    sc = min(max_w / w, max_h / h)
    s.shapes.add_picture(path, Inches(left + (max_w - w * sc) / 2), Inches(top + (max_h - h * sc) / 2), Inches(w * sc), Inches(h * sc))


def box(s, text, left, top, width, height, size=18):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    sh.fill.solid()
    sh.fill.fore_color.rgb = LIGHT
    sh.line.color.rgb = MUTED
    tf = sh.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].text = text
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    tf.paragraphs[0].font.size, tf.paragraphs[0].font.color.rgb = Pt(size), MUTED


def table(s, rows, widths, top=1.55, size=18, height=None):
    t = s.shapes.add_table(len(rows), len(rows[0]), Inches(0.6), Inches(top), Inches(sum(widths)),
                           Inches(height or 0.45 * len(rows))).table
    for j, w in enumerate(widths):
        t.columns[j].width = Inches(w)
    for i, r in enumerate(rows):
        for j, val in enumerate(r):
            t.cell(i, j).text = val
            for p in t.cell(i, j).text_frame.paragraphs:
                p.font.size, p.font.bold = Pt(size), i == 0
                p.space_before = p.space_after = Pt(0)
            t.cell(i, j).margin_top = t.cell(i, j).margin_bottom = Inches(0.02)


# 1 title
s = P.slides.add_slide(BLANK)
tf = s.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.7), Inches(2.4)).text_frame
tf.word_wrap = True
tf.paragraphs[0].text = "Physiology-Aligned Multiresolution Analysis for Patient-Independent Heart-Murmur Detection"
tf.paragraphs[0].font.size, tf.paragraphs[0].font.bold, tf.paragraphs[0].font.color.rgb = Pt(36), True, INK
p = tf.add_paragraph()
p.text = "A Comparative Study of Wavelet and Fourier Features"
p.font.size, p.font.color.rgb = Pt(26), MUTED
bullets(s, ["Parth Mishra — National Institute of Technology Warangal",
            "EC418 Biomedical Instrumentation and Signal Processing · Mini Research Project",
            "  [Roll number and course instructor to be added]",
            "  Prepared with substantial AI assistance (Claude, via Claude Code); see the report's statement"],
        top=4.2, height=2.8, size=22)
s.notes_slide.notes_text_frame.text = ("This project asks which part of a classical heart-sound pipeline carries the murmur information: "
                                       "the cardiac-cycle alignment or the time-frequency transform. Every number on these slides is "
                                       "generated from saved result tables, and every hypothesis was registered before it was tested.")

# 2 clinical motivation
s = slide("Clinical motivation",
          "A murmur is an extra sound from turbulent blood flow, heard between the first and second heart sounds. Detecting it "
          "is a screening step that leads to echocardiography; it is not a diagnosis of the underlying disease. Automatic "
          "analysis could help where trained listeners are scarce.")
bullets(s, ["Murmurs: |extra sounds from turbulent flow|, systolic or diastolic", "Innocent or a sign of structural disease (septal defects, PDA, valve lesions)",
            "Auscultation is cheap but depends on the listener's skill", "Murmur detection = |screening|, not diagnosis (echocardiography decides)",
            "Children are recorded at several sites and visits → evaluation must be |patient-independent|"],
        top=1.5, width=6.4, size=22)
image(s, f"{RF}/fig1_raw_pcg_present_vs_absent.png", 7.1, 1.5, 5.8, 5.2)

# 3 research question
s = slide("Research question",
          "Many pipelines use cycle-aligned wavelet features and report good results. We separate the two ingredients with "
          "matched controls: the same features with and without cardiac-cycle timing, and wavelet versus Fourier bands with "
          "identical band edges. We also measure how much the expert segmentation used at test time is worth.")
bullets(s, ["Does |cardiac-cycle timing| improve patient-level murmur detection?",
            "Does the |transform| matter — wavelet packets vs matched STFT bands vs scattering?",
            "Does it survive |automatic segmentation|, without the expert labels?",
            "Is an automatic gain real, or a |shortcut| from which spans are kept?",
            "Does a CirCor model |transfer| to an independent cohort?"], size=24)

# 4 related work
s = slide("Related work (2024–2026)",
          "Recent papers report high numbers, but on different scales: different metrics, label sets, units and splits, and "
          "some cut their test data at expert annotations. Only the 2025 wavelet baseline was re-evaluated on our folds, so it "
          "is the one direct comparison.",
          cite="Values as reported by the authors (abstract or stated section); references/g3b/case_study_numbers.csv.")
table(s, [["Work", "Approach", "Expert seg. at test", "Headline (as reported)"],
          ["Vimalajeewa 2025, Sci Rep", "multiscale wavelet features", "no", "re-evaluated on our folds"],
          ["Niizumi 2024, EMBC", "pre-trained audio representations", "no", f"W.acc {v('litNiizumiWacc')}"],
          ["Nie 2024, MLSP", "2-D CNN, 5-s snippets", "yes (annotated periods)", f"accuracy {v('litNieAccuracy')} % (snippets)"],
          ["Hong 2025, IEEE JBHI", "periodic-pattern transformer", "no", f"W.acc {v('litDppatWacc')} %"],
          ["Kutlu & Çolak 2026, J Clin Med", "external zero-shot tests", "no", f"external AUROC {v('litKutluExtHetero')} / {v('litKutluExtHomo')}"],
          ["Ladislaus 2026, arXiv", "noise robustness, aggregation", "no", f"3-class W.acc {v('litLadislausWacc')}"],
          ["Nogueira 2026, arXiv", "multi-channel CNN", "yes (expert S1 onsets)", f"(Se+Sp)/2 {v('litNogueiraOverall')}"]],
      [3.4, 3.6, 2.4, 2.7], size=18, height=4.9)

# 5 dataset
s = slide("Dataset and patient independence",
          "CirCor is the PhysioNet Challenge 2022 training set with expert murmur labels and segmentation. We use Present "
          "versus Absent, verified the official download against PhysioNet's checksums, and group linked visits so that no "
          "child appears in both training and test data.",
          cite="Oliveira et al., IEEE JBHI 2022; physionet.org/content/circor-heart-sound/1.0.3; Jia et al., IEEE TBME 2024 (ZCHSound).")
bullets(s, ["CirCor DigiScope v1.0.3: |942 patients, 3,163 recordings, 4 sites, 4 kHz",
            "Task: |874 patients (179 Present, 695 Absent); Unknown excluded",
            "Linked visits share an Additional ID → |816 subjects; folds never split a subject",
            "5 repeats × 5 subject-grouped folds, |fixed before any model ran",
            "Official PhysioNet download verified file by file (10,434 checksums)",
            f"External: ZCHSound, |{v('zchNClean')} clean recordings ({v('zchNchdClean')} CHD) + {v('zchNNoisy')} noisy; CHD labels, not murmurs"], size=21)

# 6 methodology
s = slide("Methodology",
          "Each recording is filtered, segmented into cardiac-cycle phases, described by band features, summarised per patient "
          "and classified in WEKA. The controls keep the features but destroy the timing, use whole-signal features, or give "
          "the baseline the same span information, so every contrast isolates one ingredient.")
image(s, f"{RR}/slide_pipeline.png", 0.4, 1.5, 12.5, 3.7)
bullets(s, ["Preprocessing: |2 kHz, band-pass 25–900 Hz, z-score",
            "Segmentation: |expert labels (oracle), envelope peaks, LR-HSMM, or HSMM + learned span detector"], top=5.35, height=1.5, size=20)

# 7 wavelet decomposition
s = slide("Wavelet decomposition",
          "The wavelet packet transform splits both branches of the filter bank, giving 16 equal 62.5 Hz bands at depth 4. "
          "The STFT bands use the same edges and time grid, so the comparison changes only the transform. Real wavelet filters "
          "leak between bands, as the packet responses show.",
          cite="Daubechies db4 (8 taps), depth J = 4, fs = 2 kHz; band width fs / 2^(J+1). Scattering: Andén & Mallat 2014; kymatio.")
bullets(s, ["DWT: |low-pass h, high-pass g, downsample by 2 at each level",
            "WPD: |split both branches → 16 bands of 62.5 Hz, 8 ms grid",
            "STFT bands: |same 16 edges, same grid, same dimension",
            "Features per band: |systolic vs S1, systolic vs diastolic energy"], top=1.45, width=12.1, height=2.4, size=21)
image(s, f"{RF}/fig4_dwt_decomposition.png", 0.5, 3.9, 6.0, 3.0)
image(s, f"{RF}/fig5_wpd_time_frequency.png", 6.8, 3.9, 6.0, 3.0)

# 8 WEKA
s = slide("WEKA classifier evaluation",
          "All models are WEKA 3.9.7 classifiers called in batch from Python on subject-grouped folds; WEKA's own cross-validation "
          "is never used because it would split a child's recordings. One fold run on the WEKA command line reproduces the "
          "pipeline's predictions exactly; the Explorer screenshots will show the same fold.",
          cite="Frank, Hall & Witten, The WEKA Workbench, 2016; Breiman 2001 (RF); Platt 1998 (SMO).")
bullets(s, ["RandomForest -I 300 -K 0 -S 1 (primary); SMO, linear PolyKernel (secondary)",
            "Thresholds from inner training folds only (specificity ≥ 0.80)",
            f"Demo fold: {v('wekaDemoNTrain')} train / {v('wekaDemoNTest')} test patients; CLI = pipeline exactly; "
            f"at 0.5: TN {v('wekaDemoTN')}, FP {v('wekaDemoFP')}, FN {v('wekaDemoFN')}, TP {v('wekaDemoTP')}"], top=1.45, height=2.4, size=20)
for i, (fn, lab) in enumerate([("1_preprocess.png", "Preprocess tab"), ("2_rf_options.png", "RF options dialog"),
                               ("3_result_buffer.png", "Result buffer + confusion matrix")]):
    path = f"weka_demo/screenshots/{fn}"
    if os.path.exists(path):
        image(s, path, 0.6 + i * 4.1, 4.1, 3.9, 2.7)
    else:
        box(s, f"Screenshot placeholder\n{lab}\n{fn}", 0.6 + i * 4.1, 4.1, 3.9, 2.7)

# 9 baseline comparison
s = slide("Baseline comparison",
          "The 2025 Scientific Reports wavelet baseline was re-implemented from the authors' code and evaluated on our patient-"
          "level folds. On this common scale the aligned and automatic models are better. Published numbers from other papers "
          "are not on this scale and are not compared directly.",
          cite="Vimalajeewa, Lee & Vidaković, Sci Rep 15:10315, 2025. AUROC of repeat-averaged patient scores, 95 % CI.")
table(s, [["Model (WEKA RF, 874 patients)", "AUROC [95 % CI]"],
          ["2025 wavelet baseline (B1)", v("v1B1RFAucFull")],
          ["Whole-signal STFT bands", v("decWholeStftFull")],
          ["Expert-aligned STFT (oracle)", v("decExpertStftFull")],
          ["Automatic spans, STFT (E4)", v("decEFourStftFull")]], [6.5, 4.0], size=20, height=2.6)
bullets(s, [f"Aligned WPD − B1 (H4b): |{v('hypH4bWpdFull')}", f"E4 − B1 (H9b): |{rng('famRangeH9b')}"], top=4.6, height=1.8, size=21)

# 10 timing
s = slide("Timing-alignment experiments",
          "Keeping the features but destroying their alignment to the cardiac cycle lowers AUROC in all 25 folds. The diagnostics "
          "split the gain exactly: within-span timing explains it, and the expert's choice of spans adds nothing positive. Rows "
          "with a star use the expert annotation at test time and are diagnostics only.")
image(s, f"{RR}/slide_decomposition.png", 0.3, 1.45, 8.4, 5.4)
bullets(s, [f"Timing destroyed (perm): |{v('hFivePermRange')}", f"Shifted labels: |{v('hFiveShiftRange')}",
            f"Within-span timing: |{v('decDiffWithinSpanTimingRange')}", f"Span choice without timing: {v('decDiffSpanSelNoTimingRange')}"],
        left=8.8, top=1.6, width=4.3, height=5.2, size=19)

# 11 automatic segmentation
s = slide("Automatic segmentation",
          "Simple envelope peaks and an HSMM on whole recordings were not good enough. With a learned span detector the point "
          "estimate comes within the pre-registered margin of expert labels, but the confidence interval is too wide to call it "
          "non-inferior. This is the honest reading of the result.",
          cite="LR-HSMM: Springer, Tarassenko & Clifford, IEEE TBME 2016 (Python port, GPL-3.0).")
bullets(s, [f"E1 envelope peaks, E1b LR-HSMM: |targets missed (e.g. H8a {v('hypH8aWpdFull')})",
            f"HSMM timing inside expert spans is close to expert: {v('decDiffHsmmTimingErrorRange')}",
            f"What is missing is span selection: {v('decDiffSpanSelHsmmRange')}",
            f"E4 learned span detector vs expert (H9a): |{v('hypH9aWpdFull')}",
            "  point-estimate rule (≥ −0.03) met; statistical non-inferiority NOT established"], size=21)

# 12 wavelet vs Fourier
s = slide("Wavelet vs Fourier findings",
          "Under every labelling, wavelet packets were at most equal to matched STFT bands, and scattering did not help either. "
          "This is still a meaningful result: because bands, grid and features were matched, it shows that the transform itself "
          "brings no advantage here, so any wavelet advantage reported elsewhere needs another explanation.")
image(s, f"{RR}/slide_wavelet.png", 0.3, 1.45, 8.2, 5.4)
bullets(s, [f"WPD − STFT: |{v('wvsAllRange')}", "Equivalence within ±0.03: |not shown",
            f"Scattering − STFT (H11a): {v('gFourH11aAScat1AStftFull')}", "STFT is at least as good — |no wavelet superiority"],
        left=8.6, top=1.6, width=4.5, height=5.2, size=19)

# 13 external test
s = slide(f"External test: ZCHSound AUROC {v('hTenA')}",
          "A whole-signal STFT model trained once on all CirCor patients was applied unchanged to an independent paediatric cohort. "
          "The subtypes follow the pre-registered physiological order, and the silent defect PFO is at chance. ZCHSound is "
          "labelled with CHD diagnoses, not murmurs, and the noisy neonatal subset is at chance.",
          cite="ZCHSound: Jia et al., IEEE TBME 71(8), 2024 (CC BY 4.0).")
image(s, f"{RR}/slide_external.png", 0.3, 1.45, 8.6, 4.6)
bullets(s, [f"Clean set (H10a): |{v('hTenAFull')}", f"VSD − PFO (H10b): {v('hTenBvsdPfoFull')}",
            f"Noisy neonates: {v('zchWStftRFAllNoisy')} (chance)", "CHD label, not murmur"], left=9.0, top=1.6, width=4.1, height=4.8, size=19)

# 14 shortcut + limitations
s = slide("Shortcut learning and limitations",
          "The automatic pipeline's span fraction predicts the class by itself. Once a whole-signal model gets the same information, "
          "the remaining advantage has a confidence interval that includes zero, so it is not established. The other limitations "
          "are listed plainly.")
bullets(s, [f"Span fraction alone predicts the class: AUROC {v('eFourShortcutPredFrac')}",
            f"Shortcut-free E4q vs STFT + span info (H12a): |{v('gTwoH12aStftFull')} — not established",
            "Cross-validation on one public cohort; no hidden test set; no clinical study",
            "Only 179 Present patients; Unknown class excluded",
            "Expert-label results are oracle-conditioned; sequential design after earlier results",
            "SMO unstable on scattering features; B0's external score is an artefact"], size=21)

# 15 conclusions
s = slide("Conclusions",
          "Timing matters, the transform does not, and a simple whole-signal model transfers to an independent CHD cohort with the "
          "expected physiological ordering. The automatic pipeline is close to expert labels on the point estimate but not shown "
          "to beat a shortcut-aware baseline. Nothing here supports clinical deployment.")
bullets(s, ["Cardiac-cycle |timing| carries the murmur information",
            "The |transform| (WPD, STFT, scattering) does not change the result",
            "Automatic segmentation needs |span selection|; non-inferiority not established",
            "Transfer to paediatric CHD with the |physiological order|",
            "Every hypothesis, failure and erratum is reported"], size=24)

# 16 future work
s = slide("Future work",
          "The next steps follow from the limitations: noisy recordings, learned models under the same strict protocol, and an "
          "external test with murmur labels. The literature audit was searched and screened but not coded, so it stays future work.")
bullets(s, ["Signal-quality gating for noisy (neonatal) recordings",
            "Learned models evaluated under the same subject-grouped, shortcut-controlled protocol",
            "A murmur-labelled external cohort and the CirCor hidden test set",
            "A systematic audit of test-time expert segmentation in published CirCor work",
            "Reproduce: scripts/p5_verify_integrity.py, docs/REPRODUCIBILITY.md"], size=23)

# 17 backup
s = slide("Backup: every pre-registered hypothesis (RF)",
          "This table summarises every pre-registered hypothesis family with the range over representations and the recorded "
          "verdict; the full table with confidence intervals is in the report appendix. SUPPORTED for H9a means only that the "
          "point-estimate rule was met.")
table(s, [["Hyp.", "Contrast", "Δ AUROC", "Verdict"],
          ["H2a/b", "WPD − STFT (aligned / whole)", f"{v('famRangeH2aMin')} / {v('famRangeH2bMin')}", "not supported"],
          ["H3", "aligned − whole WPD", v("famRangeH3Min"), "supported"],
          ["H4a/b", "aligned WPD − B0 / B1", f"{v('famRangeH4aMin')} / {v('famRangeH4bMin')}", "no / supported"],
          ["H5a–d", "expert − timing destroyed", f"{v('hFivePermMin')} to {v('hFiveShiftMax')}", "supported"],
          ["H6, H8", "naive / HSMM − expert", f"{v('famRangeH8cAMin')} to {v('famRangeH6aMax')}", "not supported"],
          ["H9a/b", "E4 − expert / B1", f"{rng('famRangeH9a')} / {rng('famRangeH9b')}", "rule met / supported"],
          ["H9c", "E4 − W_stft", rng("famRangeH9c"), "stft, wpds only"],
          ["S1", "E4 − (W_stft + SC)", v("sOneWstftRange"), "possibly shortcut"],
          ["H10a/b", "ZCHSound AUROC / subtype order", f"{v('hTenA')} / {v('hTenBvsdPfo')}", "supported"],
          ["H11a–c", "scattering − STFT", f"{v('gFourH11aAScat1AStft')} to {v('gFourH11cAScat12AScat1')}", "not supported"],
          ["H12a", "E4q − (W_stft + SC_q)", rng("famRangeH12a"), "possibly shortcut"],
          ["H12b", "E4q − random windows", rng("famRangeH12b"), "supported"]], [1.4, 4.2, 4.2, 2.3], top=1.45, size=18, height=5.4)

os.makedirs("submission", exist_ok=True)
P.save(OUT)
print(f"slides: {len(P.slides)} -> {OUT}")
